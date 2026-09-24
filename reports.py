"""
reports.py — Archivos Excel legibles para gerencia (openpyxl): encabezado del hospital, fecha de corte,
columnas en español, anchos, colores del semáforo y totales. No depende de Streamlit.

  * purchase_order_xlsx : orden de compra de los ítems urgentes (reemplaza al CSV con punto y coma).
  * executive_report_xlsx: informe ejecutivo (resumen para el administrador, situación actual, pronósticos
                           en lenguaje claro, compras urgentes, urgencias por turno). Un área sin pronóstico dice "Estimación no
                           disponible para esta área" y el resto del informe se genera igual.
Todas las cifras llegan ya calculadas desde los datos; aquí solo se da formato.
"""
from __future__ import annotations

import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HOSPITAL = "Hospital Susana López de Valencia E.S.E."
NAVY, GREEN, LIME = "2E3378", "507643", "8AB94F"
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_SOFT = PatternFill("solid", fgColor="EEF4E8")
FILL_TONE = {"danger": "FDE2E1", "warn": "FEF3C7", "ok": "DCFCE7", "neutral": "F1F5F9"}
THIN = Side(style="thin", color="D5DBE3")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
LOGO = Path(__file__).resolve().parent / "assets" / "logo_hslv_icono.png"


def _header(ws, title: str, subtitle: str, width_cols: int) -> int:
    """Logo + nombre del hospital + título. Devuelve la primera fila libre."""
    try:
        from openpyxl.drawing.image import Image
        if LOGO.exists():
            img = Image(str(LOGO))
            img.height, img.width = 54, 54 * img.width / img.height
            ws.add_image(img, "A1")
    except Exception:  # Pillow ausente: el informe sale igual, sin logo
        pass
    last = get_column_letter(max(width_cols, 2))
    for row, (text, font) in enumerate([(HOSPITAL, Font(bold=True, size=13, color=NAVY)),
                                        (title, Font(bold=True, size=12, color=GREEN)),
                                        (subtitle, Font(size=10, color="5B6472"))], start=1):
        ws.merge_cells(f"B{row}:{last}{row}")
        cell = ws[f"B{row}"]
        cell.value, cell.font = text, font
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 11)
    for c in range(1, width_cols + 1):
        ws.cell(row=4, column=c).border = Border(bottom=Side(style="medium", color=LIME))
    return 6


def _table(ws, row: int, headers: list[str], rows: list[list], widths: list[float],
           tones: list[str | None] | None = None, number_cols: dict[int, str] | None = None) -> int:
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font, cell.fill, cell.border = Font(bold=True, color="FFFFFF"), FILL_HEAD, BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = max(ws.column_dimensions[get_column_letter(c)].width or 0, w)
    for i, values in enumerate(rows, start=1):
        tone = (tones or [None] * len(rows))[i - 1]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=row + i, column=c, value=v)
            cell.border, cell.alignment = BOX, WRAP
            if number_cols and c in number_cols:
                cell.number_format = number_cols[c]
            if tone:
                cell.fill = PatternFill("solid", fgColor=FILL_TONE.get(tone, "FFFFFF"))
    return row + len(rows) + 2


def _section(ws, row: int, text: str, span: int) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font, cell.fill = Font(bold=True, size=11, color=NAVY), FILL_SOFT
    return row + 1


def _paragraph(ws, row: int, text: str, span: int, height: float = 30) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.alignment = WRAP
    ws.row_dimensions[row].height = height
    return row + 1


def _bytes(wb: Workbook) -> bytes:
    for ws in wb.worksheets:  # al imprimir o pasar a PDF: horizontal y todas las columnas en una página de ancho
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = "landscape", 1, 0
        ws.sheet_view.showGridLines = False
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Orden de compra
# ---------------------------------------------------------------------------
def _order_rows(items: list[dict]) -> list[list]:
    return [[r["codigo"], str(r["nombre"]).capitalize(), r.get("tipo_item") or "",
             int(r["disponible"]), round(float(r["consumo_diario_promedio"] or 0), 1),
             int(r["orden_sugerida_15d"])] for r in items]


ORDER_HEAD = ["Código", "Medicamento o insumo", "Tipo", "Quedan (und.)", "Uso por día (und.)",
              "Cantidad a pedir (und.)"]
ORDER_W = [15, 58, 18, 13, 14, 16]
ORDER_FMT = {4: "#,##0", 5: "#,##0.0", 6: "#,##0"}


def purchase_order_xlsx(items: list[dict], cutoff: str, generated_by: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Orden de compra"
    sub = f"Fecha de corte {cutoff} · {len(items)} ítems que se agotan en menos de 5 días"
    row = _header(ws, "Orden de compra · ítems urgentes", sub + (f" · generado por {generated_by}" if generated_by else ""),
                  len(ORDER_HEAD))
    row = _paragraph(ws, row, "Cantidad a pedir = lo necesario para cubrir 15 días de uso promedio (últimos 30 días "
                     "de dispensación real), menos lo que queda. Existencias iniciales simuladas.", len(ORDER_HEAD))
    start = row + 1
    end = _table(ws, start, ORDER_HEAD, _order_rows(items), ORDER_W, number_cols=ORDER_FMT)
    total_row = end - 1
    ws.cell(row=total_row, column=5, value="Total").font = Font(bold=True)
    cell = ws.cell(row=total_row, column=6, value=f"=SUM(F{start + 1}:F{start + len(items)})" if items else 0)
    cell.font, cell.number_format = Font(bold=True), "#,##0"
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(ORDER_HEAD))}{start + len(items)}"
    return _bytes(wb)


# ---------------------------------------------------------------------------
# Informe ejecutivo
# ---------------------------------------------------------------------------
ADMIN_QUESTIONS = [("what", "¿Qué está ocurriendo?"), ("when", "¿En qué periodo?"),
                   ("changes", "¿Qué cambio se espera?"), ("resources", "¿Qué recursos podrían verse afectados?"),
                   ("review", "¿Qué debería revisar el administrador?")]
ADMIN_EMPTY = {"resources": "Ninguno por encima de lo habitual.", "review": "Nada adicional.", "changes": "Sin datos."}
ADMIN_NOTE = ("Las sugerencias son puntos de revisión, no órdenes de compra ni de contratación: los modelos "
              "estiman cuántas atenciones habrá, no cuántos recursos se necesitan.")


def _admin_rows(admin: dict | None) -> list[tuple[str, str | list[str]]]:
    """[(pregunta, respuesta)] del resumen para el administrador (forecast_text.admin_summary().__dict__)."""
    if not admin:
        return []
    return [(q, admin.get(k) or ADMIN_EMPTY.get(k, "—")) for k, q in ADMIN_QUESTIONS]


def executive_report_xlsx(data: dict) -> bytes:
    """data = {
        cutoff, generated_by,
        situation: [(indicador, valor, lectura, tono)],
        summary: str, recommendations: [str],
        forecasts: [{area, headline, range_text, level, level_text, level_tone, confidence, confidence_text,
                     actions}] | [{area, unavailable: True}],
        purchases: [items de stock_semaphore ROJO],
        shifts: [(turno, ingresos, espera_promedio_min)],
    }"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    span = 3
    row = _header(ws, "Informe ejecutivo de operación y alerta temprana",
                  f"Fecha de corte {data['cutoff']}" + (f" · generado por {data['generated_by']}"
                                                       if data.get("generated_by") else ""), span)
    admin = _admin_rows(data.get("admin"))
    if admin:
        row = _section(ws, row, "Resumen para el administrador", span)
        rows = [[q, "\n".join(f"• {x}" for x in a) if isinstance(a, list) else a, ""] for q, a in admin]
        start = row
        row = _table(ws, row, ["Pregunta", "Respuesta", ""], rows, [30, 22, 64],
                     tones=[data["admin"].get("tone")] + [None] * (len(rows) - 1))
        for r in range(start + 1, start + 1 + len(rows)):  # la respuesta ocupa las dos columnas de la derecha
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            ws.row_dimensions[r].height = max(30, 15 * (str(ws.cell(row=r, column=2).value).count("\n") + 1) + 6)
        ws.cell(row=start, column=2).value = "Respuesta"
        ws.merge_cells(start_row=start, start_column=2, end_row=start, end_column=3)
        row = _paragraph(ws, row - 1, ADMIN_NOTE, span, height=28) + 1
    row = _section(ws, row, "Resumen", span)
    row = _paragraph(ws, row, data.get("summary", ""), span, height=80) + 1
    row = _section(ws, row, "Situación actual", span)
    sit = data.get("situation") or []
    row = _table(ws, row, ["Indicador", "Valor", "Qué significa"], [[a, b, c] for a, b, c, _ in sit],
                 [30, 22, 64], tones=[t for *_, t in sit])
    row = _section(ws, row, "Recomendaciones", span)
    for rec in data.get("recommendations") or []:
        row = _paragraph(ws, row, f"• {rec}", span, height=30)
    row += 1
    _paragraph(ws, row, "Fuente: extracto del HIS del reto (datos sintéticos) y microservicios predictivos del "
               "equipo. Las existencias de farmacia son simuladas; el consumo es real.", span, height=28)

    fc = wb.create_sheet("Pronósticos")
    head = ["Área", "Periodo", "Qué se espera", "Frente a lo habitual", "Por qué importa", "Qué conviene revisar",
            "Confianza", "Datos recientes"]
    row = _header(fc, "Pronósticos para el siguiente día", f"Fecha de corte {data['cutoff']}", len(head))
    rows, tones = [], []
    for f in data.get("forecasts") or []:
        if f.get("unavailable"):
            rows.append([f["area"], "", "Estimación no disponible para esta área", "", "", "", "", ""])
            tones.append("neutral")
            continue
        change = f"{f['level']}. {f.get('change_text') or ''}".strip(". ") + f". {f['level_text']}"
        rows.append([f["area"], f.get("period", ""), f"{f['headline']} {f['range_text']}", change, f.get("why", ""),
                     "\n".join(f"• {a}" for a in f.get("review", [])),
                     f"{f['confidence']}: {f['confidence_text']}", "\n".join(f"• {a}" for a in f.get("actions", []))])
        tones.append(f.get("level_tone"))
    _table(fc, row, head, rows, [15, 20, 34, 30, 32, 44, 40, 48], tones=tones)
    for r in range(row + 1, row + 1 + len(rows)):
        fc.row_dimensions[r].height = 105

    po = wb.create_sheet("Compras urgentes")
    items = data.get("purchases") or []
    row = _header(po, "Compras urgentes", f"{len(items)} ítems que se agotan en menos de 5 días", len(ORDER_HEAD))
    _table(po, row, ORDER_HEAD, _order_rows(items), ORDER_W, number_cols=ORDER_FMT)
    po.freeze_panes = po.cell(row=row + 1, column=1)

    sh = wb.create_sheet("Urgencias por turno")
    shifts = data.get("shifts") or []
    row = _header(sh, "Urgencias por turno · última semana", "Ingresos y espera promedio hasta la atención", 3)
    _table(sh, row, ["Turno", "Ingresos", "Espera promedio (min)"], [list(s) for s in shifts], [16, 12, 22],
           number_cols={2: "#,##0", 3: "0.0"})
    return _bytes(wb)


# ---------------------------------------------------------------------------
# Informe ejecutivo en PDF (mismo contenido que el Excel, pensado para imprimir y entregar a la dirección)
# ---------------------------------------------------------------------------
def executive_report_pdf(data: dict) -> bytes:
    """Recibe el mismo diccionario que executive_report_xlsx. Sin datos de pacientes: solo agregados."""
    from xml.sax.saxutils import escape as x

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    navy, green, lime, muted = (colors.HexColor("#" + c) for c in (NAVY, GREEN, LIME, "5B6472"))
    tone = {k: colors.HexColor("#" + v) for k, v in FILL_TONE.items()}
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=9.5, leading=13, alignment=TA_LEFT)
    small = ParagraphStyle("s", parent=base, fontSize=8, leading=10.5, textColor=muted)
    cell = ParagraphStyle("c", parent=base, fontSize=8.5, leading=11)
    head_cell = ParagraphStyle("hc", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    h1 = ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=14, leading=17, textColor=navy)
    h2 = ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=11.5, leading=15, textColor=navy,
                        spaceBefore=10, spaceAfter=4)
    lead = ParagraphStyle("l", parent=base, fontSize=11, leading=15)

    def p(text, style=base):
        return Paragraph(x(str(text)), style)

    def table(headers, rows, widths, tones=None, numeric=()):
        body = [[p(h, head_cell) for h in headers]]
        body += [[p(v, cell) for v in r] for r in rows]
        t = Table(body, colWidths=widths, repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), navy), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D5DBE3")),
                 ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 3),
                 ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
        for i, tn in enumerate(tones or [], start=1):
            if tn in tone:
                style.append(("BACKGROUND", (0, i), (-1, i), tone[tn]))
        t.setStyle(TableStyle(style))
        return t

    def bullets(items):
        return ListFlowable([ListItem(p(i), leftIndent=10) for i in items], bulletType="bullet", start="•",
                            leftIndent=12, bulletFontSize=8)

    story = []
    title = [p(HOSPITAL + " · Popayán, Cauca", ParagraphStyle("t0", parent=base, fontName="Helvetica-Bold",
                                                               fontSize=10, textColor=green)),
             p("Informe gerencial de proyección de servicios y planificación de recursos", h1),
             p(f"Fecha de corte de los datos: {data['cutoff']}"
               + (f" · generado por {data['generated_by']}" if data.get("generated_by") else ""), small),
             p("Uso exclusivo de la Dirección Médica y Administrativa. Contiene solo cifras agregadas: ningún "
               "nombre, documento ni identificador de paciente.", small)]
    logo = Image(str(LOGO), width=1.7 * cm, height=1.7 * cm) if LOGO.exists() else Spacer(1, 1)
    head = Table([[logo, title]], colWidths=[2.1 * cm, 15.4 * cm])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, 0), 2, lime),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story += [head, Spacer(1, 10)]
    admin = _admin_rows(data.get("admin"))
    if admin:
        q_style = ParagraphStyle("q", parent=cell, fontName="Helvetica-Bold", textColor=navy)

        def answer(a):
            return bullets(a) if isinstance(a, list) else p(a, cell)

        box = Table([[p("Resumen para el administrador", ParagraphStyle("at", parent=h2, spaceBefore=0)), ""]]
                    + [[p(q, q_style), answer(a)] for q, a in admin]
                    + [[p(ADMIN_NOTE, small), ""]], colWidths=[5.2 * cm, 12.3 * cm])
        box.setStyle(TableStyle([("SPAN", (0, 0), (1, 0)), ("SPAN", (0, -1), (1, -1)),
                                 ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7F9FC")),
                                 ("LINEBEFORE", (0, 0), (0, -1), 3, tone.get(data["admin"].get("tone"), navy)
                                  if data["admin"].get("tone") != "ok" else green),
                                 ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D5DBE3")),
                                 ("LINEBELOW", (0, 1), (-1, -2), 0.3, colors.HexColor("#E3E8EF")),
                                 ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 5),
                                 ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
        story += [box, Spacer(1, 6)]
    story += [p("Resumen", h2), p(data.get("summary", ""), lead)]

    sit = data.get("situation") or []
    story += [p("Situación actual", h2),
              table(["Indicador", "Valor", "Qué significa"], [[a, b, c] for a, b, c, _ in sit],
                    [5.2 * cm, 3.8 * cm, 8.5 * cm], tones=[t for *_, t in sit])]

    for n, f in enumerate(data.get("forecasts") or [], start=1):
        block = [p(f"{n}. {f['area']}", h2)]
        if f.get("unavailable"):
            block.append(p("Datos no disponibles temporalmente: el servicio de pronóstico de esta área no respondió "
                           "al generar el informe.", small))
            story.append(KeepTogether(block))
            continue
        if f.get("period"):
            block.append(p(f"Periodo: {f['period']}", small))
        block += [p(f"{f['headline']} {f['range_text']}", lead),
                  table(["Frente a lo habitual", "Confianza del pronóstico"],
                        [[f"{f['level']}. {f['level_text']}", f"{f['confidence']}. {f['confidence_text']}"]],
                        [8.2 * cm, 9.3 * cm], tones=[f.get("level_tone")])]
        if f.get("why"):
            block += [Spacer(1, 4), Paragraph(f"<b>Por qué importa.</b> {x(f['why'])}", base)]
        if f.get("review"):
            block += [Spacer(1, 3), p("Qué conviene revisar", ParagraphStyle("rv", parent=base, fontName="Helvetica-Bold")),
                      bullets(f["review"])]
        if f.get("actions"):
            block += [Spacer(1, 3), p("Lo que muestran los datos recientes",
                                      ParagraphStyle("dr", parent=base, fontName="Helvetica-Bold")),
                      bullets(f["actions"])]
        notes = [nt for nt in f.get("notes") or [] if "supera" not in nt.lower()]
        notes.sort(key=lambda nt: 0 if ("programacion" in nt or "esquema" in nt.lower()) else 1)  # la del dato primero
        for nt in notes[:2]:
            block.append(p(f"Nota metodológica: {nt}", small))
        if f["area"] == "Urgencias" and data.get("shifts"):
            block += [Spacer(1, 4), table(["Turno (última semana)", "Ingresos", "Espera promedio"],
                                          [[t, f"{i:,}".replace(",", "."), f"{w:.0f} min"] for t, i, w in data["shifts"]],
                                          [6 * cm, 3 * cm, 3.5 * cm])]
        story.append(KeepTogether(block))
        if f["area"] == "Farmacia" and data.get("purchases"):
            items = sorted(data["purchases"], key=lambda r: -int(r["orden_sugerida_15d"]))
            top = items[:15]
            story += [Spacer(1, 6), p(f"Compras urgentes: {len(items)} ítems se agotan en menos de 5 días. "
                                      f"Los {len(top)} de mayor volumen:", base), Spacer(1, 3),
                      table(["Medicamento o insumo", "Quedan", "Uso por día", "Pedir"],
                            [[str(r["nombre"]).capitalize(), int(r["disponible"]),
                              f"{float(r['consumo_diario_promedio'] or 0):.1f}".replace(".", ","),
                              f"{int(r['orden_sugerida_15d']):,}".replace(",", ".")] for r in top],
                            [10.5 * cm, 2 * cm, 2.4 * cm, 2.6 * cm]),
                      p("Pedir = lo necesario para 15 días de uso promedio real, menos lo que queda. La lista "
                        "completa está en la orden de compra en Excel. Existencias iniciales simuladas.", small)]

    recs = data.get("recommendations") or []
    if recs:
        story += [p("Plan de acción inmediato", h2),
                  ListFlowable([ListItem(p(r), leftIndent=14) for r in recs], bulletType="1", leftIndent=14,
                               bulletFontName="Helvetica-Bold", bulletFontSize=9.5, bulletColor=navy)]

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(muted)
        canvas.drawString(2 * cm, 1.2 * cm, "Fuente: extracto del HIS del reto (datos sintéticos) y microservicios "
                          "predictivos del equipo. Solo cifras agregadas.")
        canvas.drawRightString(letter[0] - 2 * cm, 1.2 * cm, f"Página {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=letter, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.6 * cm,
                      bottomMargin=2 * cm, title="Informe gerencial HSLV", author=HOSPITAL).build(
        story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
