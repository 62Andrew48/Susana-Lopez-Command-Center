"""Cliente de microservicios: enrutamiento, respaldo y circuit breaker (sin levantar los servicios)."""
import sys
from pathlib import Path

import pytest
import requests

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import ml_services as ml  # noqa: E402


@pytest.fixture(autouse=True)
def _reset():
    ml.reset_breakers()
    yield
    ml.reset_breakers()


@pytest.mark.parametrize("question, code", [
    ("¿Cuántos ingresos a urgencias se esperan mañana?", "urgencias"),
    ("pronóstico de cirugías para los próximos días", "quirofanos"),
    ("¿Qué demanda de medicamentos se espera en farmacia?", "farmacia"),
    ("predicción de consultas ambulatorias", "consultas"),
    ("¿Cuántas camas de UCI están ocupadas hoy?", None),
    ("¿Cuál fue la espera en urgencias la última semana?", None),
])
def test_route_only_forecast_questions(question, code):
    svc = ml.route(question)
    assert (svc.code if svc else None) == code


def test_describe_warns_when_model_does_not_beat_baseline():
    pred = {"prediccion": 112.1, "intervalo": [62.6, 161.6], "objetivo": "Ingresos diarios por urgencias",
            "fecha_objetivo": "2026-09-22", "metricas": {"mae": 14.85, "mae_baseline": 14.61, "supera_baseline": False}}
    text = ml.describe(ml.BY_CODE["urgencias"], pred)
    assert "112" in text and "Solo orientativo" in text and "no acertó mejor" in text


def test_breaker_opens_after_failure_and_skips_calls(monkeypatch):
    calls = []

    def boom(*a, **k):
        calls.append(1)
        raise requests.ConnectionError("apagado")

    monkeypatch.setattr(ml.requests, "request", boom)
    with pytest.raises(ml.ServiceUnavailable):
        ml.predict("quirofanos")
    with pytest.raises(ml.ServiceUnavailable):
        ml.predict("quirofanos")  # circuito abierto: no vuelve a llamar
    assert len(calls) == 1 and ml.is_down("quirofanos") and not ml.is_down("urgencias")


def test_assistant_falls_back_to_agent_when_service_is_down(monkeypatch):
    from agent import AgentResponse
    from ui.assistant_scope import SCOPES, answer
    monkeypatch.setattr(ml.requests, "request", lambda *a, **k: (_ for _ in ()).throw(requests.Timeout()))

    class Agent:
        def ask(self, q):
            return AgentResponse(q, "Respuesta del histórico", engine="reglas")

    resp = answer("¿Cuántos ingresos a urgencias se esperan mañana?", SCOPES["completo"], agent=Agent(),
                  forecasts=True)
    assert "no está disponible" in resp.answer and "Respuesta del histórico" in resp.answer


def test_forecasts_disabled_without_permission(monkeypatch):
    from agent import AgentResponse
    from ui.assistant_scope import SCOPES, answer
    monkeypatch.setattr(ml.requests, "request", lambda *a, **k: pytest.fail("no debe llamar al microservicio"))

    class Agent:
        def ask(self, q):
            return AgentResponse(q, "ok", engine="reglas")

    assert answer("¿Cuántos ingresos a urgencias se esperan mañana?", SCOPES["completo"], agent=Agent()).answer == "ok"


# ---------------------------------------------------------------------------
# Lenguaje claro e informes
# ---------------------------------------------------------------------------
import io  # noqa: E402

import forecast_text as ft  # noqa: E402
import reports  # noqa: E402

PRED_URG = {"prediccion": 112.11, "intervalo": [62.62, 161.6], "fecha_objetivo": "2026-09-22",
            "metricas": {"mae": 14.85, "mae_baseline": 14.61, "supera_baseline": False}}
KPI_URG = {"promedio_diario_28d_previos": 118.8, "cumplimiento_meta_triage2_pct_7d": 23.6, "meta_triage2_min": 30,
           "por_turno_7d": [{"turno": "Mañana", "ingresos": 334, "espera_promedio_min": 59.6},
                            {"turno": "Noche", "ingresos": 203, "espera_promedio_min": 57.3},
                            {"turno": "Tarde", "ingresos": 259, "espera_promedio_min": 61.9}]}


def test_interpret_speaks_plain_language_from_real_numbers():
    r = ft.interpret("urgencias", PRED_URG, KPI_URG)
    assert r.headline == "Se esperan unos 112 ingresos a urgencias el martes 22 de septiembre."
    assert r.range_text == "Lo normal sería entre 63 y 162."
    assert r.level == "Habitual" and "119" in r.level_text
    assert r.confidence == "Solo orientativo"
    assert any("mañana" in a and "42 %" in a for a in r.actions)      # 334 / 796
    assert any("24 %" in a for a in r.actions)                        # triage II
    joined = " ".join([r.headline, r.level_text, r.confidence_text, *r.actions]).lower()
    assert "mae" not in joined and "baseline" not in joined and "random" not in joined


@pytest.mark.parametrize("value, level", [(140, "Alta"), (118, "Habitual"), (90, "Baja")])
def test_interpret_level_against_usual(value, level):
    assert ft.interpret("urgencias", {**PRED_URG, "prediccion": value}, KPI_URG).level == level


def test_interpret_consultas_without_data_says_so():
    pred = {"prediccion": 0.44, "intervalo": [0, 1.4], "fecha_objetivo": "2026-09-22",
            "metricas": {"mae": 0.48, "mae_baseline": 0.46, "supera_baseline": False}}
    r = ft.interpret("consultas", pred, {"participacion_ambulatoria_28d_pct": 0.2})
    assert r.level == "Sin datos suficientes"


def test_purchase_order_xlsx_has_spanish_columns_and_total():
    from openpyxl import load_workbook
    items = [{"codigo": "A1", "nombre": "ACETAMINOFEN 500 MG", "tipo_item": "Medicamento", "disponible": 3,
              "consumo_diario_promedio": 2.5, "orden_sugerida_15d": 35},
             {"codigo": "B2", "nombre": "GASA ESTERIL", "tipo_item": "Insumo / dispositivo", "disponible": 1,
              "consumo_diario_promedio": 4.0, "orden_sugerida_15d": 59}]
    ws = load_workbook(io.BytesIO(reports.purchase_order_xlsx(items, "2026-09-21"))).active
    values = [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]
    assert "Cantidad a pedir (und.)" in values and "Acetaminofen 500 mg" in values
    assert any(isinstance(v, str) and v.startswith("=SUM(F") for v in values)


def test_executive_report_marks_unavailable_area():
    from openpyxl import load_workbook
    data = {"cutoff": "2026-09-21", "generated_by": "Admin", "summary": "Resumen.", "recommendations": ["Hacer X."],
            "situation": [("Camas", "86 %", "Poco margen", "warn")], "purchases": [], "shifts": [("Mañana", 334, 59.6)],
            "forecasts": [{"area": "Urgencias", **ft.interpret("urgencias", PRED_URG, KPI_URG).__dict__},
                          {"area": "Quirófanos", "unavailable": True}]}
    wb = load_workbook(io.BytesIO(reports.executive_report_xlsx(data)))
    assert wb.sheetnames == ["Resumen", "Pronósticos", "Compras urgentes", "Urgencias por turno"]
    cells = [c for row in wb["Pronósticos"].iter_rows(values_only=True) for c in row if c]
    assert "Estimación no disponible para esta área" in cells


def test_executive_report_pdf_builds_and_has_no_patient_fields():
    PdfReader = pytest.importorskip("pypdf").PdfReader
    data = {"cutoff": "2026-09-21", "generated_by": "Admin", "summary": "Resumen.", "recommendations": ["Hacer X."],
            "situation": [("Camas", "86 %", "Poco margen", "warn")], "shifts": [("Mañana", 334, 59.6)],
            "purchases": [{"codigo": "A1", "nombre": "ACETAMINOFEN", "disponible": 3, "consumo_diario_promedio": 2.5,
                           "orden_sugerida_15d": 35}],
            "forecasts": [{"area": "Urgencias", **ft.interpret("urgencias", PRED_URG, KPI_URG).__dict__},
                          {"area": "Farmacia", "unavailable": True}]}
    pdf = reports.executive_report_pdf(data)
    assert pdf.startswith(b"%PDF")
    text = " ".join(pg.extract_text() for pg in PdfReader(io.BytesIO(pdf)).pages)
    assert "112" in text and "Datos no disponibles temporalmente" in text and "Plan de acción" in text
