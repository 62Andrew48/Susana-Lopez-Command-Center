"""
verificar_servicios.py — Comprueba el contrato REST de los microservicios en ejecución.

    python microservicios/run_services.py          # en una terminal
    python microservicios/verificar_servicios.py   # en otra

Revisa en cada servicio: /health, /kpis, /predict (sin filtros) y /reporte.xlsx (exactamente tres hojas
con los nombres acordados), y que ninguna respuesta exponga identificadores de pacientes o ingresos.
Devuelve código de salida 1 si algo falla (útil antes de hacer merge a develop).
"""
from __future__ import annotations

import io
import json
import os
import re
import sys
import urllib.error
import urllib.request

from openpyxl import load_workbook

SERVICIOS = {"urgencias": ("URGENCIAS_URL", 5001), "quirofanos": ("QUIROFANOS_URL", 5002),
             "farmacia": ("FARMACIA_URL", 5003), "consultas": ("CONSULTAS_URL", 5004)}
HOJAS = ["Historico clasificado", "Consumo o uso", "Prediccion"]
PROHIBIDAS = {"idpaciente", "idpaciente2", "oidingreso", "consecutivoingreso", "consecutivoprogramacion",
              "oidtriage", "nombrepaciente", "documento", "numerodocumento", "fechanacimiento"}


def url_base(variable: str, puerto: int) -> str:
    valor = (os.getenv(variable) or "").strip()
    if valor.startswith("http"):
        return valor.rstrip("/")
    return f"http://127.0.0.1:{valor if valor.isdigit() else puerto}"


def pedir(url: str, cuerpo: dict | None = None) -> tuple[int, bytes]:
    datos = json.dumps(cuerpo).encode() if cuerpo is not None else None
    req = urllib.request.Request(url, data=datos, headers={"Content-Type": "application/json"} if datos else {})
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            return r.status, r.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()


def claves(obj) -> set[str]:
    if isinstance(obj, dict):
        return {re.sub(r"[^a-z0-9]", "", k.lower()) for k in obj} | set().union(*(claves(v) for v in obj.values()))
    if isinstance(obj, list):
        return set().union(*(claves(v) for v in obj)) if obj else set()
    return set()


def verificar(nombre: str, base: str) -> list[str]:
    fallos = []
    try:
        codigo, cuerpo = pedir(f"{base}/health")
    except OSError as exc:
        return [f"no responde en {base} ({exc})"]
    h = json.loads(cuerpo)
    if codigo != 200 or h.get("estado") != "ok" or h.get("servicio") != nombre:
        fallos.append(f"/health inesperado: {h}")
    if not isinstance(h.get("modelo_entrenado"), bool) or not isinstance(h.get("fecha_referencia"), str):
        fallos.append("/health: tipos de modelo_entrenado o fecha_referencia incorrectos")

    codigo, cuerpo = pedir(f"{base}/kpis")
    k = json.loads(cuerpo)
    if codigo != 200 or not isinstance(k, dict):
        fallos.append(f"/kpis devolvió {codigo}")
    elif claves(k) & PROHIBIDAS:
        fallos.append(f"/kpis expone identificadores: {claves(k) & PROHIBIDAS}")

    codigo, cuerpo = pedir(f"{base}/predict", {})
    p = json.loads(cuerpo)
    if codigo != 200:
        fallos.append(f"/predict devolvió {codigo}: {p}")
    else:
        m = p.get("metricas", {})
        if not isinstance(p.get("prediccion"), (int, float)):
            fallos.append("/predict: 'prediccion' no es numérico")
        if not (isinstance(p.get("intervalo"), list) and len(p["intervalo"]) == 2):
            fallos.append("/predict: 'intervalo' debe ser [inferior, superior]")
        if not {"mae", "rmse", "r2", "mae_baseline"} <= set(m):
            fallos.append(f"/predict: faltan métricas {({'mae', 'rmse', 'r2', 'mae_baseline'} - set(m))}")
        if not isinstance(p.get("features_usadas"), list):
            fallos.append("/predict: 'features_usadas' no es lista")
        if claves(p) & PROHIBIDAS:
            fallos.append("/predict expone identificadores")
        print(f"   predicción {p.get('fecha_objetivo')}: {p.get('prediccion')} {p.get('intervalo')} · "
              f"MAE {m.get('mae')} vs base {m.get('mae_baseline')} · supera={m.get('supera_baseline')}")

    codigo, cuerpo = pedir(f"{base}/reporte.xlsx")
    if codigo != 200:
        fallos.append(f"/reporte.xlsx devolvió {codigo}")
    else:
        libro = load_workbook(io.BytesIO(cuerpo), read_only=True)
        if libro.sheetnames != HOJAS:
            fallos.append(f"reporte: hojas {libro.sheetnames}, se esperaban {HOJAS}")
        for hoja in libro.worksheets:
            for fila in hoja.iter_rows(values_only=True):
                celdas = {re.sub(r"[^a-z0-9]", "", str(c).lower()) for c in fila if isinstance(c, str)}
                if celdas & PROHIBIDAS:
                    fallos.append(f"reporte: columna identificadora en '{hoja.title}'")
                    break
    return fallos


def main() -> int:
    total = 0
    for nombre, (variable, puerto) in SERVICIOS.items():
        base = url_base(variable, puerto)
        print(f"· {nombre} ({base})")
        fallos = verificar(nombre, base)
        total += len(fallos)
        for f in fallos:
            print(f"   ✘ {f}")
        if not fallos:
            print("   ✔ contrato completo")
    print("\nResultado:", "OK" if total == 0 else f"{total} fallo(s)")
    return 1 if total else 0


if __name__ == "__main__":
    sys.exit(main())
